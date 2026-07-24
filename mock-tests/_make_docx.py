#!/usr/bin/env python3
"""Build a print-ready Word (.docx) version of the 70-day study plan
from mock-tests/manifest.json. Output: STUDY_PLAN.docx (repo root)."""

import json
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
manifest = json.load(open(os.path.join(ROOT, "mock-tests", "manifest.json")))
tests = manifest["tests"]

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x55, 0x55, 0x55)
HDR_FILL = "1F3864"
ALT_FILL = "EEF2FA"

phase_titles = {
    1: "PHASE 1 — Foundation (Topic-wise Building)  ·  Days 1–35",
    2: "PHASE 2 — Consolidation (Multi-topic + First Full Mocks)  ·  Days 36–55",
    3: "PHASE 3 — Simulation & Revision (Full-length + Weak-area)  ·  Days 56–69",
}
phase_notes = {
    1: ("3 tests every day — Morning 30Q, Afternoon 40Q, Late 50Q. Subjects are interleaved so no "
        "subject goes cold. Difficulty starts at 'foundation' (easy-tilted) and moves to 'standard' "
        "from Day 17. Every 7th day the late session is a cumulative revision of that week's topics."),
    2: ("3 tests every day — Morning 50Q (full subject), Afternoon 60Q (full subject), Late 50Q "
        "revision or a 100Q FULL MOCK. Difficulty is 'advanced'. Full-length 100-question mocks "
        "land on Days 39, 43, 47, 51 and 55."),
    3: ("3 tests every day — Morning 100Q FULL MOCK (exam simulation), Afternoon 50Q sectional "
        "revision, Late 40Q rapid/weak-area drill. Difficulty is 'exam' / 'advanced'. Day 70 is the "
        "EXAM — no new test, only light confidence revision and rest."),
}


def set_cell_bg(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)


doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)

# Narrow margins to fit tables
for section in doc.sections:
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

# ---- Title block ----
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("70-Day Study & Mock-Test Plan")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("JKSSB Sub-Inspector (J&K Police) — Written Examination")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Medium: English   |   Exam: Day 70   |   Daily study: 7–10 hours across 3 sessions").italic = True

stat = doc.add_paragraph()
stat.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = stat.add_run(f"Total planned tests: {manifest['total_tests']}    "
                  f"Total questions: {sum(t['total_questions'] for t in tests):,}")
rs.bold = True
rs.font.color.rgb = NAVY

doc.add_paragraph()
how = doc.add_paragraph()
how.add_run("How to use each day: ").bold = True
how.add_run("Each morning, revise the topic(s)/subject(s) listed for that day, then attempt the "
            "day's mock tests in order (Morning → Afternoon → Late). Tick the Status box as you "
            "complete each test. Difficulty profiles: foundation (easy-tilted), standard (blueprint), "
            "advanced (hard-tilted), exam (full-length simulation).")

# ---- Per-phase, per-day tables ----
cols = ["#", "Session", "Subject", "Topic / Focus", "Q", "Difficulty", "Type", "Done"]
widths = [Inches(0.4), Inches(0.8), Inches(1.7), Inches(2.9), Inches(0.4),
          Inches(0.9), Inches(1.3), Inches(0.5)]

current_phase = None
days = sorted(set(t["day"] for t in tests))
for day in days:
    day_tests = [t for t in tests if t["day"] == day]
    phase = day_tests[0]["phase"]
    if phase != current_phase:
        current_phase = phase
        doc.add_page_break()
        h = doc.add_paragraph()
        hr = h.add_run(phase_titles[phase])
        hr.bold = True
        hr.font.size = Pt(14)
        hr.font.color.rgb = NAVY
        note = doc.add_paragraph()
        note.add_run(phase_notes[phase]).italic = True
        doc.add_paragraph()

    dh = doc.add_paragraph()
    dhr = dh.add_run(f"Day {day}")
    dhr.bold = True
    dhr.font.size = Pt(12)
    dhr.font.color.rgb = NAVY

    table = doc.add_table(rows=1, cols=len(cols))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, c in enumerate(cols):
        cell = hdr.cells[i]
        cell.width = widths[i]
        cell.paragraphs[0].text = ""
        run = cell.paragraphs[0].add_run(c)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)
        set_cell_bg(cell, HDR_FILL)

    for ridx, t in enumerate(day_tests):
        subj = t["subject"] if t["subject"] == "FULL" else f'{t["subject"]} — {t["subject_name"]}'
        vals = [str(t["number"]), t["session"], subj, t["topics"],
                str(t["total_questions"]), t["difficulty_profile"], t["type"], "☐"]
        row = table.add_row()
        for i, v in enumerate(vals):
            cell = row.cells[i]
            cell.width = widths[i]
            cell.paragraphs[0].text = ""
            run = cell.paragraphs[0].add_run(v)
            run.font.size = Pt(9)
            if i == 0:
                run.bold = True
            if ridx % 2 == 1:
                set_cell_bg(cell, ALT_FILL)
    doc.add_paragraph()

# ---- Exam day ----
doc.add_page_break()
eh = doc.add_paragraph()
ehr = eh.add_run("Day 70 — EXAM DAY")
ehr.bold = True
ehr.font.size = Pt(16)
ehr.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
ep = doc.add_paragraph()
ep.add_run("No new mock test. Light revision of formula sheets, current affairs and previously-marked "
           "weak points only. Reach the exam centre early, stay calm, and manage negative marking "
           "wisely (skip only if truly unsure). All the best!")

out = os.path.join(ROOT, "STUDY_PLAN.docx")
doc.save(out)
print("Saved:", out)
print("Days:", len(days), "| Tests:", len(tests))
